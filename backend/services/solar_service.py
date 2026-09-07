"""
Plantas solares GEM Utility-Scale (1 MW+) en Colombia.

Fuente: data/Global-Solar-Power-Tracker-February-2026.xlsx
Hoja: Utility-Scale (1 MW+)
Filtros fijos: Country=Colombia, Status=operating|construction.
Recorte al bioma ejecutado con el vector de regiones EE (una vez, en caché).
"""

from __future__ import annotations

import json
import logging
from functools import lru_cache
from pathlib import Path
from typing import Any

from fastapi import HTTPException

from backend.core.biomas import expandir_biomas
from backend.core.cache import cache
from backend.core.paths import PATHS
from backend.core.runtime_paths import get_install_root

logger = logging.getLogger(__name__)

_STATUSES = frozenset({"operating", "construction"})
_JSON_NAME = "solar_colombia_utility.json"
_XLSX_NAME = "Global-Solar-Power-Tracker-February-2026.xlsx"
_SHEET = "Utility-Scale (1 MW+)"


def _data_dir() -> Path:
    return get_install_root() / "data"


def _json_path() -> Path:
    return _data_dir() / _JSON_NAME


def _xlsx_path() -> Path:
    return _data_dir() / _XLSX_NAME


def _refresh_json_if_needed() -> None:
    xlsx = _xlsx_path()
    dest = _json_path()
    if dest.is_file() and (
        not xlsx.is_file() or dest.stat().st_mtime >= xlsx.stat().st_mtime
    ):
        return
    if not xlsx.is_file():
        return
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[_SHEET]
    it = ws.iter_rows(values_only=True)
    header = [str(h or "").strip() for h in next(it)]
    idx = {h: i for i, h in enumerate(header)}
    points: list[dict[str, Any]] = []
    for row in it:
        country = str(row[idx["Country/Area"]] or "").strip()
        if country.lower() != "colombia":
            continue
        status = str(row[idx["Status"]] or "").strip().lower()
        if status not in _STATUSES:
            continue
        try:
            lat = float(row[idx["Latitude"]])
            lon = float(row[idx["Longitude"]])
        except (TypeError, ValueError):
            continue
        cap = row[idx["Capacity (MW)"]]
        try:
            cap_f = float(cap) if cap not in (None, "") else None
        except (TypeError, ValueError):
            cap_f = None
        gem = str(
            row[idx["GEM phase ID"]] or row[idx["GEM location ID"]] or ""
        ).strip()
        points.append(
            {
                "id": gem or f"{row[idx['Project Name']]}-{lat}-{lon}",
                "name": str(row[idx["Project Name"]] or "").strip(),
                "phase": str(row[idx["Phase Name"]] or "").strip().strip("-"),
                "status": status,
                "mw": cap_f,
                "tech": str(row[idx["Technology Type"]] or "").strip(),
                "lat": lat,
                "lon": lon,
                "province": str(row[idx["State/Province"]] or "").strip(),
                "wiki": str(row[idx["Wiki URL"]] or "").strip(),
            }
        )
    wb.close()
    dest.write_text(
        json.dumps(
            {"source": xlsx.name, "count": len(points), "points": points},
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    logger.info("Solar GEM extract written: %s plants", len(points))


@lru_cache(maxsize=1)
def _load_points() -> list[dict[str, Any]]:
    _refresh_json_if_needed()
    path = _json_path()
    if not path.is_file():
        raise HTTPException(
            status_code=503,
            detail=(
                "Datos solares no encontrados. Coloca "
                f"{_XLSX_NAME} en data/ o genera {_JSON_NAME}."
            ),
        )
    data = json.loads(path.read_text(encoding="utf-8"))
    return list(data.get("points") or [])


def _index_path() -> Path:
    return _data_dir() / "solar_biome_index.json"


def _read_disk_index() -> dict[str, str]:
    path = _index_path()
    if not path.is_file():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return {str(k): str(v) for k, v in (data.get("index") or {}).items()}
    except Exception:
        logger.exception("No se pudo leer solar_biome_index.json")
        return {}


def _write_disk_index(index: dict[str, str]) -> None:
    _index_path().write_text(
        json.dumps({"index": index}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def _points_fc(plants: list[dict[str, Any]]):
    import ee

    return ee.FeatureCollection(
        [
            ee.Feature(
                ee.Geometry.Point([p["lon"], p["lat"]]),
                {"id": str(p["id"])},
            )
            for p in plants
        ]
    )


def _ids_in_biome_raster(biome: str, plants: list[dict[str, Any]]) -> list[str]:
    """Sample a 500 m biome raster — avoids loading detailed vectors into memory."""
    import ee

    from backend.core import ee_client  # noqa: F401

    mask = (
        ee.FeatureCollection(PATHS["region_vector"])
        .filter(ee.Filter.eq("bioma", biome))
        .reduceToImage(["id_regionC"], ee.Reducer.first())
        .gt(0)
        .selfMask()
    )
    sampled = mask.reduceRegions(
        collection=_points_fc(plants),
        reducer=ee.Reducer.first(),
        scale=500,
        tileScale=4,
    )
    kept = sampled.filter(ee.Filter.notNull(["first"]))
    return [str(x) for x in (kept.aggregate_array("id").getInfo() or [])]


def _ids_in_biome_bbox(biome: str, plants: list[dict[str, Any]]) -> list[str]:
    """Fallback: bbox of the biome (cheap, slightly looser than the outline)."""
    import ee

    from backend.core import ee_client  # noqa: F401

    geom = (
        ee.FeatureCollection(PATHS["region_vector"])
        .filter(ee.Filter.eq("bioma", biome))
        .geometry()
        .bounds(maxError=5000)
    )
    kept = _points_fc(plants).filterBounds(geom)
    return [str(x) for x in (kept.aggregate_array("id").getInfo() or [])]


def _ensure_biome_tags(internos: list[str]) -> dict[str, str]:
    """id → bioma, persistido en disco. Completa solo los biomas pedidos."""
    cache_key = "solar_biome_index_v2"
    index = cache.get(cache_key) if cache_key in cache else None
    if index is None:
        index = _read_disk_index()

    untagged = [p for p in _load_points() if str(p["id"]) not in index]
    if not untagged:
        cache[cache_key] = index
        return index

    dirty = False
    remaining = untagged
    for biome in internos:
        if not remaining:
            break
        try:
            ids = _ids_in_biome_raster(biome, remaining)
        except Exception:
            logger.warning(
                "Solar raster clip failed for %s; using bbox fallback",
                biome,
                exc_info=True,
            )
            try:
                ids = _ids_in_biome_bbox(biome, remaining)
            except Exception:
                logger.exception("Solar bbox clip also failed for %s", biome)
                continue
        hit = set(ids)
        for p in remaining:
            pid = str(p["id"])
            if pid in hit:
                index[pid] = biome
                dirty = True
        remaining = [p for p in remaining if str(p["id"]) not in index]
        logger.info("Solar clip %s: +%s tagged, %s left", biome, len(hit), len(remaining))

    if dirty:
        _write_disk_index(index)
    cache[cache_key] = index
    return index


def listar_plantas(biomas: list[str]) -> dict[str, Any]:
    internos = expandir_biomas(biomas)
    if not internos:
        raise HTTPException(status_code=400, detail="Selecciona al menos un bioma.")

    plants = _load_points()
    index = _ensure_biome_tags(internos)
    allowed = set(internos)
    points = []
    for p in plants:
        bio = index.get(str(p["id"]))
        if bio in allowed:
            points.append({**p, "bioma": bio})

    return {
        "source": "GEM Global Solar Power Tracker February 2026",
        "sheet": _SHEET,
        "statuses": sorted(_STATUSES),
        "biomas": internos,
        "total": len(points),
        "points": points,
    }
