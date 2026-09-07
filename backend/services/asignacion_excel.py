"""
Intérprete responsable por región (columna A × regionId en B).

Fuente por defecto: Excel local del avance Col4.
Opcional: Google Sheets online (AVANCE_SHEET_ID) — misma estructura de columnas.
"""

from __future__ import annotations

import logging
import os
import re
from functools import lru_cache
from typing import Any, Optional

import gspread
from openpyxl import load_workbook

from backend.core.config import CREDS

logger = logging.getLogger(__name__)

_HEADER_SKIP = re.compile(
    r"^(int[eé]rprete|responsable|asignaci[oó]n)",
    re.IGNORECASE,
)


def _default_xlsx_path() -> str:
    try:
        from backend.core.config import AVANCE_COLOMBIA_XLSX

        return AVANCE_COLOMBIA_XLSX
    except Exception:
        root = os.path.dirname(
            os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        )
        return os.path.join(
            root,
            "data",
            "COLOMBIA_COL_4_Formatos de avance detallado Colombia.xlsx",
        )


def _parse_interpreter_rows(rows: list[list[Any]]) -> dict[str, str]:
    """Columna A = intérprete, B = regionId → { '30205': 'Janner', ... }."""
    out: dict[str, str] = {}
    for row in rows:
        if not row:
            continue
        nombre = row[0] if len(row) > 0 else None
        rid = row[1] if len(row) > 1 else None
        if nombre is None or rid is None:
            continue
        nombre_s = str(nombre).strip()
        if not nombre_s or _HEADER_SKIP.match(nombre_s):
            continue
        try:
            rid_s = str(int(float(rid)))
        except (TypeError, ValueError):
            continue
        if rid_s not in out:
            out[rid_s] = nombre_s
    return out


def _resolve_avance_worksheet(spreadsheet: gspread.Spreadsheet, hoja: str, gid: int | None):
    if gid is not None:
        for ws in spreadsheet.worksheets():
            if ws.id == gid:
                return ws
        logger.warning(
            "Hoja gid=%s no encontrada en spreadsheet; probando por nombre %r",
            gid,
            hoja,
        )
    if hoja in [ws.title for ws in spreadsheet.worksheets()]:
        return spreadsheet.worksheet(hoja)
    logger.warning("Hoja %r no existe en spreadsheet online", hoja)
    return None


def _load_from_sheets() -> dict[str, str]:
    from backend.core.config import AVANCE_SHEET, AVANCE_SHEET_GID, AVANCE_SHEET_ID

    if not AVANCE_SHEET_ID:
        return {}

    try:
        client = gspread.authorize(CREDS)
        spreadsheet = client.open_by_key(AVANCE_SHEET_ID)
        ws = _resolve_avance_worksheet(spreadsheet, AVANCE_SHEET, AVANCE_SHEET_GID)
        if ws is None:
            return {}
        rows = ws.get("A:B") or []
        out = _parse_interpreter_rows(rows)
        logger.info(
            "Intérpretes cargados desde Google Sheets (%s): %s regiones",
            AVANCE_SHEET_ID,
            len(out),
        )
        return out
    except Exception:
        logger.exception("Error leyendo intérpretes desde Google Sheets")
        return {}


def _load_from_excel(ruta: Optional[str] = None) -> dict[str, str]:
    path = ruta or os.environ.get("AVANCE_COLOMBIA_XLSX") or _default_xlsx_path()
    path = os.path.normpath(path)
    if not os.path.isfile(path):
        logger.warning("Excel de avance no encontrado: %s", path)
        return {}

    try:
        from backend.core.config import AVANCE_SHEET as hoja_cfg
    except Exception:
        hoja_cfg = "MAPA GENERAL COLOMBIA"
    hoja = os.environ.get("AVANCE_SHEET", hoja_cfg)
    out: dict[str, str] = {}
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        if hoja not in wb.sheetnames:
            logger.warning("Hoja %s no existe en %s", hoja, path)
            wb.close()
            return {}
        ws = wb[hoja]
        rows = list(ws.iter_rows(min_row=1, min_col=1, max_col=2, values_only=True))
        wb.close()
        out = _parse_interpreter_rows([list(r) for r in rows])
    except Exception:
        logger.exception("Error leyendo intérpretes desde Excel")
        return {}

    logger.info("Intérpretes cargados desde Excel: %s regiones", len(out))
    return out


@lru_cache(maxsize=2)
def mapa_interprete_por_region(ruta: Optional[str] = None) -> dict[str, str]:
    """
    Lee columna A (intérprete) y B (regionId).
    Prioridad: Google Sheets (AVANCE_SHEET_ID) → Excel local.
    """
    from backend.core.config import AVANCE_SHEET_ID

    if AVANCE_SHEET_ID:
        online = _load_from_sheets()
        if online:
            return online
        logger.warning(
            "AVANCE_SHEET_ID configurado pero lectura online falló; "
            "intentando Excel local."
        )

    return _load_from_excel(ruta)


def interprete_de_region(region_id) -> Optional[str]:
    if region_id is None:
        return None
    try:
        key = str(int(float(region_id)))
    except (TypeError, ValueError):
        key = str(region_id).strip()
    return mapa_interprete_por_region().get(key)
