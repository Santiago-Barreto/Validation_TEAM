"""Extract GEM Utility-Scale Colombia plants (operating/construction) to JSON."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "data" / "Global-Solar-Power-Tracker-February-2026.xlsx"
OUT = ROOT / "data" / "solar_colombia_utility.json"
SHEET = "Utility-Scale (1 MW+)"
STATUSES = frozenset({"operating", "construction"})


def main() -> int:
    if not XLSX.is_file():
        raise SystemExit(f"Missing {XLSX}")
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb[SHEET]
    it = ws.iter_rows(values_only=True)
    header = [str(h or "").strip() for h in next(it)]
    idx = {h: i for i, h in enumerate(header)}
    out: list[dict] = []
    for row in it:
        country = str(row[idx["Country/Area"]] or "").strip()
        if country.lower() != "colombia":
            continue
        status = str(row[idx["Status"]] or "").strip().lower()
        if status not in STATUSES:
            continue
        try:
            lat = float(row[idx["Latitude"]])
            lon = float(row[idx["Longitude"]])
        except (TypeError, ValueError):
            continue
        cap = row[idx["Capacity (MW)"]]
        try:
            cap_f = float(cap) if cap not in (None, "") else None
        except (TypeError, ValueError):
            cap_f = None
        gem = str(
            row[idx["GEM phase ID"]] or row[idx["GEM location ID"]] or ""
        ).strip()
        out.append(
            {
                "id": gem or f"{row[idx['Project Name']]}-{lat}-{lon}",
                "name": str(row[idx["Project Name"]] or "").strip(),
                "phase": str(row[idx["Phase Name"]] or "").strip().strip("-"),
                "status": status,
                "mw": cap_f,
                "tech": str(row[idx["Technology Type"]] or "").strip(),
                "lat": lat,
                "lon": lon,
                "province": str(row[idx["State/Province"]] or "").strip(),
                "wiki": str(row[idx["Wiki URL"]] or "").strip(),
            }
        )
    wb.close()
    OUT.write_text(
        json.dumps(
            {"source": XLSX.name, "count": len(out), "points": out},
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(f"wrote {OUT} n={len(out)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
